from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

import openpyxl

from ecommerce_manager.app.services.import_export_service import ImportExportService
from ecommerce_manager.app.services.purchases_service import PurchasesService
from ecommerce_manager.app.services.restock_service import RestockService
from ecommerce_manager.app.services.sales_service import SalesService
from tests.conftest import make_listing


def test_purchase_receiving_increases_inventory(session):
    _, product, _ = make_listing(session, initial_units=1)
    supplier = product.supplier
    if supplier is None:
        from ecommerce_manager.app.services.products_service import ProductsService

        supplier = ProductsService(session).create_supplier(name="Supplier")
    purchases = PurchasesService(session)
    po = purchases.create_purchase_order(
        supplier_id=supplier.id,
        order_date=date(2026, 5, 1),
        subtotal=Decimal("100.00"),
    )
    line = purchases.add_line(
        purchase_order_id=po.id,
        base_product_id=product.id,
        ordered_physical_units=10,
        supplier_lot_cost=Decimal("20.00"),
        supplier_lot_units=10,
    )
    batch = purchases.receive_line(purchase_order_line_id=line.id, received_date=date(2026, 5, 5))
    assert batch.remaining_physical_units == 10
    assert po.status == "received"


def test_restock_zero_sales_does_not_crash(session):
    _, _, listing = make_listing(session, units_per_listing=2, initial_units=20)
    recommendation = RestockService(session).recommendation_for_listing(
        listing.id,
        as_of_date=date(2026, 5, 31),
        velocity_days=30,
    )
    assert recommendation.average_daily_listing_sales == Decimal("0.0000")
    assert recommendation.days_of_coverage is None
    assert recommendation.urgency_status == "healthy"


def test_restock_recommends_before_stockout(session):
    _, _, listing = make_listing(session, units_per_listing=2, initial_units=100)
    sales = SalesService(session)
    as_of = date(2026, 5, 31)
    for i in range(10):
        sales.add_manual_sale(
            sale_date=as_of - timedelta(days=i),
            sku=listing.sku,
            quantity_listings_sold=1,
            gross_revenue=Decimal("10.00"),
            net_received=Decimal("8.00"),
        )
    recommendation = RestockService(session).recommendation_for_listing(
        listing.id,
        as_of_date=as_of,
        velocity_days=10,
        supplier_lead_time_days=5,
        shipping_days=5,
        safety_stock_days=5,
        target_stock_days=60,
    )
    assert recommendation.average_daily_physical_consumption == Decimal("2.0000")
    assert recommendation.total_replenishment_days == 15
    assert recommendation.reorder_point_physical_units == 30
    assert recommendation.estimated_stockout_date == as_of + timedelta(days=40)
    assert recommendation.suggested_reorder_date == as_of + timedelta(days=25)


def test_workbook_preview_flags_missing_net_received(tmp_path, session):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro de Ventas"
    ws.append(["Fecha", "SKU", "Producto", "Unidades Vendidas", "Ingreso Bruto ($)", "Ingreso Neto", "Mes", "Año"])
    ws.append([date(2026, 5, 1), "SKU-1", "Product", 1, 10, None, "mayo", 2026])
    ws.add_table(openpyxl.worksheet.table.Table(displayName="Ventas", ref="A1:H2"))
    for title, table_name, headers in [
        ("Catálogo y Costos", "Costos_Catalogo", ["SKU", "ASIN", "Nombre del Producto", "Costo Total ($)", "Precio Venta ($)", "Stock FBA Inicial", "BuyBox Mensual (%)"]),
        ("Análisis de productos", "TB_Analisis", ["Nombre del producto", "Categoría del producto", "ASIN", "SKU"]),
        ("Simulador de compraventa", "Tabla_Simulador", ["ASIN", "Nombre"]),
    ]:
        sheet = wb.create_sheet(title)
        sheet.append(headers)
        sheet.add_table(openpyxl.worksheet.table.Table(displayName=table_name, ref=f"A1:{chr(64 + len(headers))}1"))
    wb.save(path)

    preview = ImportExportService(session).preview_workbook(path)
    assert preview.sales_rows == 1
    assert "missing net received" in preview.warnings[0]

